import pytest
import io
import csv
import openpyxl
import docx
from pypdf import PdfWriter

from app.services.document_service import DocumentProcessingService
from app.schemas.document import DocumentType, DocumentExtractionStatus


@pytest.fixture
def doc_service():
    return DocumentProcessingService()


def test_coa_full_extraction_from_text(doc_service):
    raw_text = """
    CERTIFICATE OF ANALYSIS (COA)
    Product Name: Paracetamol USP Grade
    Batch Number: B-2026-9901
    Test Date: 2026-02-10
    Manufacturing Date: 2026-01-15
    Expiry Date: 2028-01-14
    Testing Facility: Eurofins BioPharma Testing Lab

    TEST PARAMETERS:
    Assay (Purity): 99.45%
    Moisture (Water Content): 0.32%
    Heavy Metals: < 5.0 ppm
    Total Aerobic Microbial Count (TAMC): < 10 cfu/g
    Escherichia coli: Absent
    """

    data, missing = doc_service.extract_coa(raw_text)

    assert data.batch_number == "B-2026-9901"
    assert data.material == "Paracetamol USP Grade"
    assert data.purity == 99.45
    assert data.moisture == 0.32
    assert data.heavy_metals == 5.0
    assert data.microbial_parameters["TAMC_cfu_g"] == 10.0
    assert data.microbial_parameters["E_coli"] == "Absent"
    assert data.test_date == "2026-02-10"
    assert data.manufacturing_date == "2026-01-15"
    assert data.expiry_date == "2028-01-14"
    assert "Eurofins" in data.laboratory
    assert len(missing) == 0


def test_coa_missing_critical_purity_and_batch_number(doc_service):
    # COA lacking purity and batch number
    raw_text = """
    CERTIFICATE OF ANALYSIS
    Product Name: Paracetamol
    Moisture: 0.4%
    Laboratory: Quality Lab
    """

    data, missing = doc_service.extract_coa(raw_text)

    assert data.purity is None
    assert data.batch_number is None
    assert "purity" in missing
    assert "batch_number" in missing

    # Test processing workflow flags as NEEDS_REVIEW
    res = doc_service.process_document(
        file_source=raw_text.encode("utf-8"),
        filename="test_coa.txt",
        document_type=DocumentType.COA,
    )
    assert res.status == DocumentExtractionStatus.NEEDS_REVIEW
    assert "purity" in res.missing_fields
    assert "batch_number" in res.missing_fields


def test_sds_extraction(doc_service):
    sds_text = """
    SAFETY DATA SHEET (SDS)
    Product Name: Paracetamol Pure
    CAS Number: 103-90-2
    Revision Date: 2025-11-20

    SECTION 2: HAZARD IDENTIFICATION
    H302 - Harmful if swallowed
    H319 - Causes serious eye irritation

    SECTION 7: HANDLING AND STORAGE
    Storage: Store in a cool, dry place. Keep container tightly closed below 25C.
    """

    data, missing = doc_service.extract_sds(sds_text)

    assert data.product_name == "Paracetamol Pure"
    assert data.cas_number == "103-90-2"
    assert len(data.hazard_statements) >= 2
    assert "Store in a cool" in data.storage_requirements
    assert data.revision_date == "2025-11-20"


def test_gmp_extraction(doc_service):
    gmp_text = """
    GOOD MANUFACTURING PRACTICE CERTIFICATE
    Certificate Number: GMP-EU-2025-88412
    Issuing Authority: European Medicines Agency (EMA)
    Scope of Certification: Synthesis of active pharmaceutical ingredients
    Valid Through: 2027-12-31
    """

    data, missing = doc_service.extract_gmp(gmp_text)

    assert data.certificate_number == "GMP-EU-2025-88412"
    assert "European Medicines Agency" in data.issuing_body
    assert data.valid_until == "2027-12-31"
    assert data.is_valid is True
    assert len(missing) == 0


def test_csv_extraction(doc_service):
    csv_buffer = io.StringIO()
    writer = csv.writer(csv_buffer)
    writer.writerow(["Parameter", "Specification", "Result"])
    writer.writerow(["Batch Number", "-", "B-CSV-100"])
    writer.writerow(["Purity", ">=99.0%", "99.8%"])
    writer.writerow(["Moisture", "<=0.5%", "0.15%"])
    content = csv_buffer.getvalue().encode("utf-8")

    text = doc_service.extract_raw_text(content, "test.csv")
    assert "B-CSV-100" in text
    assert "99.8%" in text

    res = doc_service.process_document(content, "test.csv", DocumentType.COA)
    assert res.extracted_data.get("purity") == 99.8
    assert res.extracted_data.get("batch_number") == "B-CSV-100"


def test_xlsx_extraction(doc_service):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "LabResults"
    ws.append(["Field", "Value"])
    ws.append(["Batch Number", "B-XLS-555"])
    ws.append(["Purity", "99.2%"])
    ws.append(["Moisture", "0.22%"])

    stream = io.BytesIO()
    wb.save(stream)
    content = stream.getvalue()

    text = doc_service.extract_raw_text(content, "results.xlsx")
    assert "B-XLS-555" in text
    assert "99.2%" in text


def test_docx_extraction(doc_service):
    doc = docx.Document()
    doc.add_heading("Certificate of Analysis", 0)
    doc.add_paragraph("Batch Number: B-DOCX-777")
    doc.add_paragraph("Purity: 99.6%")
    doc.add_paragraph("Moisture: 0.18%")

    stream = io.BytesIO()
    doc.save(stream)
    content = stream.getvalue()

    text = doc_service.extract_raw_text(content, "certificate.docx")
    assert "B-DOCX-777" in text
    assert "99.6%" in text


def test_pdf_extraction(doc_service):
    # Create a minimal valid PDF using pypdf
    writer = PdfWriter()
    writer.add_blank_page(width=200, height=200)
    stream = io.BytesIO()
    writer.write(stream)
    content = stream.getvalue()

    # Empty page should extract empty string
    text = doc_service.extract_raw_text(content, "empty.pdf")
    assert isinstance(text, str)


def test_empty_document_handling(doc_service):
    # Empty file must be marked as NEEDS_REVIEW without crashing
    res = doc_service.process_document(b"", "empty.txt", DocumentType.COA)
    assert res.status == DocumentExtractionStatus.NEEDS_REVIEW
    assert "Document is empty" in res.error


def test_malformed_corrupt_file_handling(doc_service):
    # Corrupt PDF binary stream
    corrupt_bytes = b"NOT_A_REAL_PDF_DATA_JUST_CORRUPTED_BYTES"
    res = doc_service.process_document(corrupt_bytes, "corrupt.pdf", DocumentType.COA)
    assert res.status == DocumentExtractionStatus.NEEDS_REVIEW
    assert res.error is not None
    assert "Extraction failure" in res.error


def test_unsupported_file_extension(doc_service):
    res = doc_service.process_document(b"hello", "file.unsupported", DocumentType.COA)
    assert res.status == DocumentExtractionStatus.NEEDS_REVIEW
    assert "Unsupported file format" in res.error
